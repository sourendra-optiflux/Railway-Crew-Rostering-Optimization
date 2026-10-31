"""OptiFlux database schema: the 12 source tables and their conversion to an :class:`Instance`.

Input data is supplied as the tables of the OptiFlux database, one CSV file per
table, named after the table (``roster_df.csv``, ``activities_df.csv``, ...).
Column names are used exactly as in the database, including spaces and case.
:data:`DB_TABLES` lists every table and its columns. Columns marked in
:data:`ASSUMED_COLUMNS` could not be read from the schema screenshots and are
a documented guess (see ``docs/db_schema.md``). They can be renamed in one
place here.

:func:`build_instance` turns the tables into the optimiser's internal
:class:`~railway_rostering.domain.Instance`, and :func:`model_config_from_constraints`
turns ``constraints_df`` into a :class:`~railway_rostering.config.ModelConfig`.
The conversion works like this. ``docs/db_schema.md`` covers each rule and the
reason for it.

* **Who is optimised.** Crew whose ``Classification`` is ``Relief`` or
  ``Excess``. Every other classification (e.g. ``Permanent``) is regular crew.
  Their roster rows count as coverage, but they are not re-planned.
* **What needs covering.** ``activity_demand_df`` gives the number of crew
  required per activity and shift *every day* of the planning horizon.
  Inactive rows and ``excluded_activities`` are skipped. A demand slot
  (activity, date, shift) is covered by matching ``roster_df`` rows (same
  activity, same start date and time). A regular-crew row counts only if that
  crew member can still legally work it (medical, restrictions, onsite status,
  competencies, currency, role). Otherwise it is a *disruption* and the slot is
  open again. Whatever stays uncovered becomes the open demand of a duty.
* **Existing Relief/Excess work.** A Relief/Excess roster row with
  ``Is Locked = Yes`` (or on an activity with ``Can Be Moved = No``) is a fixed
  assignment. An unlocked row is released, so the optimiser may re-plan that
  crew member and re-cover the slot. ``Activity Type = Standby`` rows mark the
  days a Relief/Excess crew member is rostered to be available, and their
  ``Shift Start Time`` is the rostered start time.
* **Eligibility** for a (crew, duty) pair is checked per duty date. It needs
  medical fitness, no restriction on that date, every competency the
  ``compliance_matrix_df`` requires for the activity (``status`` valid and not
  expired by the duty's end), currency (last performed within
  ``Currency Threshold`` days), a ``Role`` allowed by ``Roles/Job``, and no
  open mentorship on ``Mentor Required`` activities.
* **Priority** comes from ``activity_priority``. Row 1 is the highest
  priority, and the weight of rank *r* out of *N* is ``N + 1 - r``.
* **Rules and horizon** come from ``constraints_df`` (see
  :data:`CONSTRAINT_FIELDS`).
"""

from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Iterable, Optional

from .config import ModelConfig, model_config_from_dict
from .domain import Crew, CrewType, Duty, Instance, ShiftType

# ---------------------------------------------------------------------------
# Schema
# ---------------------------------------------------------------------------

DB_TABLES: dict[str, list[str]] = {
    "roster_df": [
        "Assignment ID", "User Id", "Start Date", "End Date", "Shift Start Time", "Shift End Time",
        "CRT Activity ID", "Activity Name", "Is Locked", "Activity Head", "Activity Type", "Activity Group",
    ],
    "activities_df": [
        "Activity ID", "CRT Activity ID", "Activity Name", "Activity Group", "Activity Head", "Activity Type",
        "TEA Required", "Can Be Moved", "Start Depot", "End Depot", "Location", "Mentor Required", "Roles/Job",
        "Currency Threshold",
    ],
    "activity_demand_df": ["Demand ID", "Shift", "Demand", "Is Active", "Activity ID", "Activity Name", "Activity Head"],
    "crew_restrictions": [
        "Restriction ID", "CM Id", "Restriction Type", "Activity Head", "Effective From", "Effective To", "Reason",
    ],
    "crew_competencies_df": ["CM Id", "Competency", "Competency LMS Code", "Expiry Date", "status"],
    "compliance_matrix_df": ["Activity ID", "Activity Name", "Competency LMS Code", "Competency", "Activity Head"],
    "crew_master_data_df": [
        "CM Id", "User Id", "Email", "First Name", "Last Name", "Mentor Id", "Classification", "Depot", "Role",
        "Is Mentor", "Medical", "Position",
    ],
    "Crew_Onsite_App": ["Record ID", "CM Id", "Date", "Location", "Status"],
    "crew_currency_df": ["Currency ID", "CM Id", "Activity Head", "Last Performed Date", "Last CRT Activity ID"],
    "constraints_df": ["Constraint", "Value", "Status"],
    "excluded_activities": ["Activity ID", "Activity Name"],
    "activity_priority": ["Priority", "Activity Head"],
}  # fmt: skip

# Columns not visible in the schema screenshots: (table, column) -> note.
ASSUMED_COLUMNS: dict[tuple[str, str], str] = {
    ("crew_restrictions", "Restriction ID"): "columns 1-4 not visible",
    ("crew_restrictions", "CM Id"): "columns 1-4 not visible",
    ("crew_restrictions", "Restriction Type"): "columns 1-4 not visible",
    ("crew_restrictions", "Activity Head"): "columns 1-4 not visible",
    ("Crew_Onsite_App", "*"): "no column visible; whole table assumed",
    ("crew_currency_df", "*"): "table name and columns 1-2 not visible (3-5 are)",
    ("excluded_activities", "*"): "no column visible; whole table assumed",
    ("activity_priority", "*"): "only the ordered values were visible; header assumed",
}

REQUIRED_TABLES = ("roster_df", "activities_df", "activity_demand_df", "crew_master_data_df", "constraints_df")
DATE_FMT = "%Y-%m-%d"
TIME_FMT = "%H:%M"

STANDBY = "standby"
OPTIMISED_CLASSIFICATIONS = {"relief": CrewType.RELIEF, "excess": CrewType.EXCESS}
VALID_COMPETENCY_STATUS = {"valid", "active", "current"}
ABSENT_ONSITE_STATUS = {"absent", "sick", "off site - unavailable"}
UNFIT_MEDICAL = {"unfit", "expired", "suspended"}

# Restriction types
RESTRICT_NO_DUTY = "no duty"
RESTRICT_ACTIVITY = "activity restriction"
RESTRICT_NO_NIGHT = "no night shift"
RESTRICTION_TYPES = (RESTRICT_NO_DUTY, RESTRICT_ACTIVITY, RESTRICT_NO_NIGHT)

# Reasons a crew member may not work an activity on a date.
R_MEDICAL = "MEDICAL_UNFIT"
R_UNAVAILABLE = "UNAVAILABLE"
R_ROLE = "ROLE_NOT_PERMITTED"
R_COMPETENCY = "COMPETENCY_MISSING_OR_EXPIRED"
R_CURRENCY = "CURRENCY_LAPSED"
R_MENTOR = "MENTOR_REQUIRED"
R_ACTIVITY = "ACTIVITY_RESTRICTED"
R_NIGHT = "NIGHT_RESTRICTED"


class DBDataError(ValueError):
    """Raised when the database tables cannot be converted. ``errors`` lists every problem."""

    def __init__(self, errors: list[str]):
        self.errors = errors
        preview = "\n  - ".join(errors[:20])
        more = f"\n  ... and {len(errors) - 20} more" if len(errors) > 20 else ""
        super().__init__(f"{len(errors)} input error(s):\n  - {preview}{more}")


# ---------------------------------------------------------------------------
# Reading / writing tables
# ---------------------------------------------------------------------------

Tables = dict[str, list[dict[str, str]]]


def is_db_dir(directory: str | Path) -> bool:
    p = Path(directory)
    return is_excel(p) or (p / "roster_df.csv").exists()


def is_excel(path: str | Path) -> bool:
    p = Path(path)
    return p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm")


# ---------------------------------------------------------------------------
# Excel workbook: one sheet per table (same names and headers as the CSVs)
# ---------------------------------------------------------------------------


def _cell_text(value) -> str:
    """Excel turns typed dates/times into Python objects; convert them back to the CSV formats."""

    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.strftime(DATE_FMT)
        return value.strftime(f"{DATE_FMT} {TIME_FMT}")
    if isinstance(value, date):
        return value.strftime(DATE_FMT)
    if isinstance(value, time):
        return value.strftime(TIME_FMT)
    if isinstance(value, timedelta):  # a time typed as a duration
        minutes = int(value.total_seconds() // 60)
        return f"{minutes // 60:02d}:{minutes % 60:02d}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_excel_tables(path: str | Path) -> Tables:
    """Read a workbook with one sheet per table (see :func:`write_excel_template`)."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise DBDataError(["reading Excel needs openpyxl: pip install openpyxl"]) from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    tables: Tables = {}
    errors: list[str] = []
    for name, columns in DB_TABLES.items():
        if name not in wb.sheetnames:
            if name in REQUIRED_TABLES:
                errors.append(f"{Path(path).name}: missing required sheet {name!r}")
            tables[name] = []
            continue
        rows = wb[name].iter_rows(values_only=True)
        header = [_cell_text(h) for h in next(rows, ())]
        while header and not header[-1]:
            header.pop()
        missing = [c for c in columns if c not in header]
        if missing:
            errors.append(f"sheet {name}: missing columns {missing}")
        out = []
        for values in rows:
            texts = [_cell_text(v) for v in values]
            if not any(texts):
                continue  # skip blank rows
            out.append({h: (texts[i] if i < len(texts) else "") for i, h in enumerate(header) if h})
        tables[name] = out
    wb.close()
    if errors:
        raise DBDataError(errors)
    return tables


def write_excel_template(path: str | Path, tables: Optional[Tables] = None) -> Path:
    """Write a workbook with one sheet per table and the exact schema headers.

    With ``tables=None`` the sheets are blank (headers only) for loading real
    data. Columns in :data:`ASSUMED_COLUMNS` are highlighted, and every header
    carries a comment saying how it is used.
    """

    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    known = PatternFill("solid", fgColor="1F4E78")
    assumed = PatternFill("solid", fgColor="C65911")
    wb = Workbook()
    wb.remove(wb.active)
    for name, columns in DB_TABLES.items():
        ws = wb.create_sheet(name)
        ws.append(columns)
        whole_table_assumed = (name, "*") in ASSUMED_COLUMNS
        for i, col in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=i)
            is_assumed = whole_table_assumed or (name, col) in ASSUMED_COLUMNS
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = assumed if is_assumed else known
            note = COLUMN_NOTES.get((name, col), "")
            if is_assumed:
                note = (note + "\n" if note else "") + "ASSUMED column name (not visible in the schema screenshots)."
            if note:
                cell.comment = Comment(note, "railway_rostering")
            ws.column_dimensions[get_column_letter(i)].width = max(14, len(col) + 4)
            # Text format so Excel keeps IDs, dates and times exactly as typed.
            for r in range(2, 1001):
                ws.cell(row=r, column=i).number_format = "@"
        ws.freeze_panes = "A2"
        for row in (tables or {}).get(name, []):
            ws.append([row.get(c, "") for c in columns])
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# How each column is read (shown as header comments in the Excel template).
COLUMN_NOTES: dict[tuple[str, str], str] = {
    ("roster_df", "Start Date"): "YYYY-MM-DD, date the shift starts",
    ("roster_df", "End Date"): "YYYY-MM-DD, next day for a night shift",
    ("roster_df", "Shift Start Time"): "HH:MM",
    ("roster_df", "Shift End Time"): "HH:MM",
    ("roster_df", "Is Locked"): "Yes/No. Relief/Excess rows: Yes = keep, No = may be re-planned",
    ("roster_df", "Activity Type"): "'Standby' = availability day of a Relief/Excess crew member",
    ("activities_df", "Can Be Moved"): "Yes/No. No = Relief/Excess assignments on it are always kept",
    ("activities_df", "Start Depot"): "Duty start location",
    ("activities_df", "End Depot"): "Duty end location",
    ("activities_df", "Mentor Required"): "Yes/No. Yes = crew with a Mentor Id may not work it",
    ("activities_df", "Roles/Job"): "Allowed crew roles separated by ';' (blank = any)",
    ("activities_df", "Currency Threshold"): "Days; blank = no currency rule",
    ("activities_df", "Activity Type"): "Use 'Standby' for the Relief/Excess availability activities",
    ("activity_demand_df", "Shift"): "HH:MM-HH:MM, e.g. 05:00-17:00 (end <= start crosses midnight)",
    ("activity_demand_df", "Demand"): "Crew required on that shift, every day of the horizon",
    ("activity_demand_df", "Is Active"): "Yes/No",
    ("crew_restrictions", "Restriction Type"): "No Duty | Activity Restriction | No Night Shift",
    ("crew_restrictions", "Effective From"): "YYYY-MM-DD (blank = open)",
    ("crew_restrictions", "Effective To"): "YYYY-MM-DD (blank = open)",
    ("crew_competencies_df", "Expiry Date"): "YYYY-MM-DD (blank = no expiry)",
    ("crew_competencies_df", "status"): "Valid / Active / Current count as held",
    ("crew_master_data_df", "Classification"): "Relief / Excess are optimised; anything else is regular crew",
    ("crew_master_data_df", "Depot"): "Base location (return-to-base rules)",
    ("crew_master_data_df", "Medical"): "Unfit / Expired / Suspended = may not work",
    ("Crew_Onsite_App", "Status"): "Absent / Sick = unavailable that day",
    ("crew_currency_df", "Last Performed Date"): "YYYY-MM-DD",
    ("constraints_df", "Constraint"): "See docs/db_schema.md section 2.10 for recognised names",
    ("constraints_df", "Status"): "Active / Inactive",
    ("activity_priority", "Priority"): "1 = highest",
}


def read_tables(directory: str | Path) -> Tables:
    """Read every table CSV present in ``directory``. Missing optional tables are empty."""

    base = Path(directory)
    tables: Tables = {}
    errors: list[str] = []
    for name, columns in DB_TABLES.items():
        path = base / f"{name}.csv"
        if not path.exists():
            if name in REQUIRED_TABLES:
                errors.append(f"missing required table {path}")
            tables[name] = []
            continue
        with path.open(newline="", encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            header = [h.strip() for h in (reader.fieldnames or [])]
            missing = [c for c in columns if c not in header]
            if missing:
                errors.append(f"{path.name}: missing columns {missing}")
            reader.fieldnames = header
            tables[name] = [{k: (v or "").strip() for k, v in row.items() if k is not None} for row in reader]
    if errors:
        raise DBDataError(errors)
    return tables


def write_tables(tables: Tables, directory: str | Path, metadata: Optional[dict] = None) -> Path:
    """Write ``tables`` as ``<table>.csv`` files (exact schema column order)."""

    out = Path(directory)
    out.mkdir(parents=True, exist_ok=True)
    for name, columns in DB_TABLES.items():
        with (out / f"{name}.csv").open("w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=columns, lineterminator="\n")
            w.writeheader()
            for row in tables.get(name, []):
                w.writerow({c: row.get(c, "") for c in columns})
    if metadata is not None:
        (out / "instance.json").write_text(json.dumps(metadata, indent=2, sort_keys=True, default=str) + "\n")
    return out


# ---------------------------------------------------------------------------
# Field parsing
# ---------------------------------------------------------------------------


def _yes(value: str) -> bool:
    return value.strip().lower() in {"y", "yes", "true", "1", "t"}


def _date(value: str) -> date:
    return datetime.strptime(value.strip()[:10], DATE_FMT).date()


def _opt_date(value: str) -> Optional[date]:
    return None if not value.strip() else _date(value)


def _time(value: str) -> time:
    return datetime.strptime(value.strip()[:5], TIME_FMT).time()


def _list(value: str) -> list[str]:
    return [p.strip() for p in value.replace(",", ";").split(";") if p.strip()]


def parse_shift(value: str) -> tuple[time, time]:
    """``"05:00-17:00"`` -> (05:00, 17:00). An end at or before the start crosses midnight."""

    a, sep, b = value.partition("-")
    if not sep:
        raise ValueError(f"Shift must look like 'HH:MM-HH:MM', got {value!r}")
    return _time(a), _time(b)


def slot_datetimes(day: date, start: time, end: time) -> tuple[datetime, datetime]:
    s = datetime.combine(day, start)
    e = datetime.combine(day, end)
    if e <= s:
        e += timedelta(days=1)
    return s, e


# ---------------------------------------------------------------------------
# constraints_df -> ModelConfig
# ---------------------------------------------------------------------------

# Constraint name (case-insensitive) -> description of how it maps. The parser
# below implements each entry; this dict is the documentation and the list of
# recognised names.
CONSTRAINT_FIELDS: dict[str, str] = {
    "planning horizon start": "first day of the horizon (YYYY-MM-DD); required",
    "planning horizon days": "number of days in the horizon; required",
    "ordinary rest hours": "ordinary_rest_hours; Inactive -> 0",
    "transition rest hours": "transition_rest_hours; Inactive -> transition_rest_enabled = false",
    "day shift window": "'HH:MM-HH:MM' start-time window of a DAY shift (day_shift_start_from/until)",
    "max duties per day": "max_duties_per_day; Inactive -> no cap",
    "relief work/off days": "'14/14' default Relief swing if the roster does not show one",
    "excess work/off days": "'6/6' default Excess swing if the roster does not show one",
    "relief allowed start offsets": "minutes from rostered time, 'ANY' or '-60;0;60'",
    "excess allowed start offsets": "minutes from rostered time, e.g. '-60;0;60' (lift-up/layback)",
    "relief time change policy": "disabled | window_only | persist_after_change; Inactive -> disabled",
    "excess time change policy": "disabled | window_only | persist_after_change; Inactive -> disabled",
    "enforce swing work days": "Status Active/Inactive",
    "base return at swing end": "Status Active/Inactive",
    "base return on transition": "Status Active/Inactive",
    "base return at horizon end": "Status Active/Inactive",
    "demand mode": "partial | all_or_nothing",
    "objective order": "levels separated by ';', e.g. 'coverage;priority;time_changes'",
    "positioning": "must be Inactive: the schema has no travel table",
}


@dataclass
class ParsedConstraints:
    horizon_start: Optional[date]
    horizon_days: Optional[int]
    overrides: dict
    warnings: list[str]


def parse_constraints(rows: list[dict[str, str]], errors: list[str]) -> ParsedConstraints:
    base = ModelConfig()
    ov: dict = {}
    warnings: list[str] = []
    h_start = h_days = None
    offsets = dict(base.allowed_offsets_minutes)
    policies = dict(base.time_change_policy)
    work_off = dict(base.default_work_off_days)
    known_fields = set(base.to_dict())
    for n, row in enumerate(rows, start=2):
        name = row.get("Constraint", "").strip()
        key = name.lower()
        value = row.get("Value", "").strip()
        status = row.get("Status", "").strip().lower() or "active"
        active = status in {"active", "yes", "y", "enabled", "on", "true"}
        try:
            if key == "planning horizon start":
                h_start = _date(value)
            elif key == "planning horizon days":
                h_days = int(value)
                if h_days < 1:
                    raise ValueError("must be >= 1")
            elif key == "ordinary rest hours":
                ov["ordinary_rest_hours"] = float(value) if active else 0.0
            elif key == "transition rest hours":
                ov["transition_rest_enabled"] = active
                if value:
                    ov["transition_rest_hours"] = float(value)
            elif key == "day shift window":
                a, b = parse_shift(value)
                ov["day_shift_start_from"] = a.strftime(TIME_FMT)
                ov["day_shift_start_until"] = b.strftime(TIME_FMT)
            elif key == "max duties per day":
                ov["max_duties_per_day"] = int(value) if active else None
            elif key in ("relief work/off days", "excess work/off days"):
                w, _, o = value.partition("/")
                work_off["RELIEF" if key.startswith("relief") else "EXCESS"] = [int(w), int(o)]
            elif key in ("relief allowed start offsets", "excess allowed start offsets"):
                ct = "RELIEF" if key.startswith("relief") else "EXCESS"
                if not active:
                    offsets[ct] = [0]
                elif value.upper() in ("", "ANY"):
                    offsets[ct] = None
                else:
                    offsets[ct] = sorted({int(v) for v in _list(value)} | {0})
            elif key in ("relief time change policy", "excess time change policy"):
                ct = "RELIEF" if key.startswith("relief") else "EXCESS"
                policies[ct] = value if active else "disabled"
            elif key == "enforce swing work days":
                ov["enforce_swing_work_days"] = active
            elif key == "base return at swing end":
                ov["base_return_at_swing_end"] = active
            elif key == "base return on transition":
                ov["base_return_on_transition"] = active
            elif key == "base return at horizon end":
                ov["base_return_at_horizon_end"] = active
            elif key == "demand mode":
                ov["demand_mode"] = value
            elif key == "objective order":
                ov["objective_levels"] = _list(value)
            elif key == "positioning":
                if active:
                    raise ValueError("positioning needs travel times, which the schema does not provide")
                ov["positioning_enabled"] = False
            elif name in known_fields:  # escape hatch: a raw ModelConfig field name
                if active:
                    try:
                        ov[name] = json.loads(value)
                    except json.JSONDecodeError:
                        ov[name] = value
            else:
                warnings.append(f"constraints_df: unknown constraint {name!r} ignored")
        except ValueError as exc:
            errors.append(f"constraints_df line {n} ({name}): {exc}")
    ov["allowed_offsets_minutes"] = offsets
    ov["time_change_policy"] = policies
    ov["default_work_off_days"] = work_off
    return ParsedConstraints(h_start, h_days, ov, warnings)


def model_config_from_constraints(
    rows: list[dict[str, str]], priority_ranks: Optional[dict[str, int]] = None
) -> ModelConfig:
    errors: list[str] = []
    pc = parse_constraints(rows, errors)
    if errors:
        raise DBDataError(errors)
    data = {**ModelConfig().to_dict(), **pc.overrides}
    if priority_ranks:
        data.update(priority_config(priority_ranks))
    return model_config_from_dict(data)


def priority_config(ranks: dict[str, int]) -> dict:
    """Weights for the priority objective: rank r of N -> N + 1 - r; unranked -> 0."""

    n = max(ranks.values(), default=0)
    return {"priority_weights": {str(r): n + 1 - r for r in range(1, n + 1)}, "default_priority_weight": 0}


# ---------------------------------------------------------------------------
# Records
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class ActivityRec:
    activity_id: str
    crt_id: str
    name: str
    group: str
    head: str
    type: str
    tea_required: bool
    can_be_moved: bool
    start_depot: str
    end_depot: str
    location: str
    mentor_required: bool
    roles: Optional[frozenset[str]]
    currency_days: Optional[int]

    @property
    def is_standby(self) -> bool:
        return self.type.lower() == STANDBY


@dataclass(frozen=True)
class CrewRec:
    cm_id: str
    user_id: str
    classification: str
    depot: str
    role: str
    mentor_id: str
    is_mentor: bool
    medical: str
    position: str

    @property
    def crew_type(self) -> Optional[CrewType]:
        return OPTIMISED_CLASSIFICATIONS.get(self.classification.strip().lower())

    @property
    def medically_fit(self) -> bool:
        return self.medical.strip().lower() not in UNFIT_MEDICAL


@dataclass(frozen=True)
class RosterRec:
    assignment_id: str
    cm_id: str
    start: datetime
    end: datetime
    activity: ActivityRec
    locked: bool


class _Rules:
    """Per-crew, per-activity, per-date legality checks shared by regular and Relief/Excess crew."""

    def __init__(self, tables: Tables, activities: dict[str, ActivityRec], crews: dict[str, CrewRec],
                 cfg: ModelConfig, errors: list[str]):  # fmt: skip
        self.activities = activities
        self.crews = crews
        self.cfg = cfg
        self.comps: dict[str, dict[str, tuple[Optional[date], bool]]] = defaultdict(dict)
        for n, r in enumerate(tables.get("crew_competencies_df", []), start=2):
            try:
                code = r["Competency LMS Code"] or r["Competency"]
                valid = r.get("status", "").lower() in VALID_COMPETENCY_STATUS
                expiry = _opt_date(r.get("Expiry Date", ""))
                prev = self.comps[r["CM Id"]].get(code)
                # Keep the best record if a competency appears twice.
                if prev is None or (valid and (not prev[1] or (prev[0] or date.max) < (expiry or date.max))):
                    self.comps[r["CM Id"]][code] = (expiry, valid)
            except (ValueError, KeyError) as exc:
                errors.append(f"crew_competencies_df line {n}: {exc}")
        self.required: dict[str, set[str]] = defaultdict(set)
        for n, r in enumerate(tables.get("compliance_matrix_df", []), start=2):
            aid = r.get("Activity ID", "")
            if aid not in activities:
                errors.append(f"compliance_matrix_df line {n}: unknown Activity ID {aid!r}")
                continue
            self.required[aid].add(r["Competency LMS Code"] or r["Competency"])
        self.currency: dict[tuple[str, str], date] = {}
        for n, r in enumerate(tables.get("crew_currency_df", []), start=2):
            try:
                d = _date(r["Last Performed Date"])
                key = (r["CM Id"], r["Activity Head"])
                self.currency[key] = max(d, self.currency.get(key, d))
            except (ValueError, KeyError) as exc:
                errors.append(f"crew_currency_df line {n}: {exc}")
        self.restrictions: dict[str, list[tuple[str, str, date, date]]] = defaultdict(list)
        for n, r in enumerate(tables.get("crew_restrictions", []), start=2):
            try:
                rtype = r["Restriction Type"].strip().lower()
                if rtype not in RESTRICTION_TYPES:
                    raise ValueError(f"Restriction Type {r['Restriction Type']!r} not in {RESTRICTION_TYPES}")
                lo = _opt_date(r.get("Effective From", "")) or date.min
                hi = _opt_date(r.get("Effective To", "")) or date.max
                self.restrictions[r["CM Id"]].append((rtype, r.get("Activity Head", ""), lo, hi))
            except (ValueError, KeyError) as exc:
                errors.append(f"crew_restrictions line {n}: {exc}")
        self.onsite: dict[str, list[tuple[date, str, str]]] = defaultdict(list)
        for n, r in enumerate(tables.get("Crew_Onsite_App", []), start=2):
            try:
                self.onsite[r["CM Id"]].append((_date(r["Date"]), r.get("Location", ""), r.get("Status", "")))
            except (ValueError, KeyError) as exc:
                errors.append(f"Crew_Onsite_App line {n}: {exc}")
        for v in self.onsite.values():
            v.sort()
        self._static: dict[tuple[str, str], tuple[Optional[str], date]] = {}

    # -- availability --------------------------------------------------------
    def unavailable(self, cm: str, day: date) -> bool:
        for rtype, _, lo, hi in self.restrictions.get(cm, ()):
            if rtype == RESTRICT_NO_DUTY and lo <= day <= hi:
                return True
        for d, _, status in self.onsite.get(cm, ()):
            if d == day and status.strip().lower() in ABSENT_ONSITE_STATUS:
                return True
        return False

    def _static_check(self, cm: str, aid: str) -> tuple[Optional[str], date]:
        """Date-independent part: (reason or None, last day the competencies stay valid)."""

        key = (cm, aid)
        if key in self._static:
            return self._static[key]
        crew, act = self.crews[cm], self.activities[aid]
        out: tuple[Optional[str], date]
        if not crew.medically_fit:
            out = (R_MEDICAL, date.min)
        elif act.roles is not None and crew.role not in act.roles:
            out = (R_ROLE, date.min)
        elif act.mentor_required and crew.mentor_id:
            out = (R_MENTOR, date.min)
        else:
            valid_until = date.max
            reason = None
            held = self.comps.get(cm, {})
            for code in self.required.get(aid, ()):
                rec = held.get(code)
                if rec is None or not rec[1]:
                    reason = R_COMPETENCY
                    break
                if rec[0] is not None:
                    valid_until = min(valid_until, rec[0])
            out = (reason, valid_until)
        self._static[key] = out
        return out

    def reason(self, cm: str, aid: str, start: datetime, end: datetime) -> Optional[str]:
        """``None`` if crew ``cm`` may work activity ``aid`` from ``start`` to ``end``."""

        reason, valid_until = self._static_check(cm, aid)
        if reason:
            return reason
        day = start.date()
        if end.date() > valid_until:
            return R_COMPETENCY
        if self.unavailable(cm, day):
            return R_UNAVAILABLE
        act = self.activities[aid]
        if act.currency_days:
            last = self.currency.get((cm, act.head))
            if last is None or (day - last).days > act.currency_days:
                return R_CURRENCY
        night = _is_night(start, self.cfg)
        for rtype, head, lo, hi in self.restrictions.get(cm, ()):
            if not lo <= day <= hi:
                continue
            if rtype == RESTRICT_ACTIVITY and (not head or head == act.head):
                return R_ACTIVITY
            if rtype == RESTRICT_NO_NIGHT and night:
                return R_NIGHT
        return None


def _is_night(start: datetime, cfg: ModelConfig) -> bool:
    from .constraints.time_rest import classify_start

    return classify_start(start, cfg) == ShiftType.NIGHT


def _shift_of(start: datetime, cfg: ModelConfig) -> ShiftType:
    return ShiftType.NIGHT if _is_night(start, cfg) else ShiftType.DAY


def _runs(days: Iterable[date]) -> list[list[date]]:
    runs: list[list[date]] = []
    for d in sorted(set(days)):
        if runs and (d - runs[-1][-1]).days == 1:
            runs[-1].append(d)
        else:
            runs.append([d])
    return runs


# ---------------------------------------------------------------------------
# Conversion
# ---------------------------------------------------------------------------


def build_instance(tables: Tables, name: str = "db_instance", metadata: Optional[dict] = None) -> Instance:
    """Convert database tables into an :class:`Instance`.

    The derived :class:`ModelConfig` (from ``constraints_df`` and
    ``activity_priority``) is stored in ``instance.metadata["model_config"]``.
    A data-preparation report (coverage by regular crew, disruptions,
    released assignments, crew excluded from optimisation) is stored in
    ``instance.metadata["data_preparation"]``.
    """

    errors: list[str] = []
    for t in REQUIRED_TABLES:
        if not tables.get(t):
            errors.append(f"table {t} is empty")
    if errors:
        raise DBDataError(errors)

    # ---- priorities and constraints ----------------------------------------
    ranks: dict[str, int] = {}
    for n, r in enumerate(tables.get("activity_priority", []), start=2):
        head = r.get("Activity Head", "")
        try:
            rank = int(r["Priority"]) if r.get("Priority") else n - 1
        except ValueError:
            errors.append(f"activity_priority line {n}: Priority {r.get('Priority')!r} is not an integer")
            continue
        if head in ranks:
            errors.append(f"activity_priority line {n}: duplicate Activity Head {head!r}")
        ranks[head] = rank
    pc = parse_constraints(tables["constraints_df"], errors)
    if pc.horizon_start is None or pc.horizon_days is None:
        errors.append("constraints_df must define 'Planning Horizon Start' and 'Planning Horizon Days'")
    if errors:
        raise DBDataError(errors)
    cfg_dict = {**ModelConfig().to_dict(), **pc.overrides, **priority_config(ranks)}
    try:
        cfg = model_config_from_dict(cfg_dict)
    except ValueError as exc:
        raise DBDataError([f"constraints_df: {exc}"]) from exc
    h0 = datetime.combine(pc.horizon_start, time(0, 0))
    h1 = h0 + timedelta(days=pc.horizon_days)
    horizon_days = [pc.horizon_start + timedelta(days=k) for k in range(pc.horizon_days)]

    # ---- activities ----------------------------------------------------------
    activities: dict[str, ActivityRec] = {}
    crt_to_id: dict[str, str] = {}
    for n, r in enumerate(tables["activities_df"], start=2):
        try:
            thr = r.get("Currency Threshold", "")
            roles = _list(r.get("Roles/Job", ""))
            a = ActivityRec(
                activity_id=r["Activity ID"], crt_id=r["CRT Activity ID"], name=r["Activity Name"],
                group=r.get("Activity Group", ""), head=r.get("Activity Head", ""), type=r.get("Activity Type", ""),
                tea_required=_yes(r.get("TEA Required", "")), can_be_moved=_yes(r.get("Can Be Moved", "yes")),
                start_depot=r.get("Start Depot", "") or r.get("Location", ""),
                end_depot=r.get("End Depot", "") or r.get("Start Depot", "") or r.get("Location", ""),
                location=r.get("Location", ""), mentor_required=_yes(r.get("Mentor Required", "")),
                roles=frozenset(roles) if roles else None, currency_days=int(float(thr)) if thr else None,
            )  # fmt: skip
        except (ValueError, KeyError) as exc:
            errors.append(f"activities_df line {n}: {exc}")
            continue
        if a.activity_id in activities:
            errors.append(f"activities_df line {n}: duplicate Activity ID {a.activity_id!r}")
        if not a.is_standby and not (a.start_depot and a.end_depot):
            errors.append(f"activities_df line {n}: activity {a.activity_id} has no Start/End Depot")
        activities[a.activity_id] = a
        if a.crt_id:
            crt_to_id[a.crt_id] = a.activity_id
    by_name = {a.name: a.activity_id for a in activities.values()}

    excluded: set[str] = set()
    for r in tables.get("excluded_activities", []):
        aid = r.get("Activity ID") or by_name.get(r.get("Activity Name", ""), "")
        if aid:
            excluded.add(aid)

    # ---- crew ------------------------------------------------------------
    crews: dict[str, CrewRec] = {}
    user_to_cm: dict[str, str] = {}
    for n, r in enumerate(tables["crew_master_data_df"], start=2):
        try:
            c = CrewRec(
                cm_id=r["CM Id"], user_id=r["User Id"], classification=r.get("Classification", ""),
                depot=r.get("Depot", ""), role=r.get("Role", ""), mentor_id=r.get("Mentor Id", ""),
                is_mentor=_yes(r.get("Is Mentor", "")), medical=r.get("Medical", ""), position=r.get("Position", ""),
            )  # fmt: skip
        except KeyError as exc:
            errors.append(f"crew_master_data_df line {n}: missing {exc}")
            continue
        if c.cm_id in crews:
            errors.append(f"crew_master_data_df line {n}: duplicate CM Id {c.cm_id!r}")
        crews[c.cm_id] = c
        user_to_cm[c.user_id] = c.cm_id

    rules = _Rules(tables, activities, crews, cfg, errors)

    # ---- roster ------------------------------------------------------------
    roster: list[RosterRec] = []
    for n, r in enumerate(tables["roster_df"], start=2):
        try:
            cm = user_to_cm.get(r["User Id"])
            if cm is None:
                raise ValueError(f"unknown User Id {r['User Id']!r}")
            aid = crt_to_id.get(r["CRT Activity ID"]) or by_name.get(r.get("Activity Name", ""))
            if aid is None:
                raise ValueError(f"unknown CRT Activity ID {r['CRT Activity ID']!r}")
            start = datetime.combine(_date(r["Start Date"]), _time(r["Shift Start Time"]))
            end = datetime.combine(_date(r["End Date"] or r["Start Date"]), _time(r["Shift End Time"]))
            if end <= start:
                raise ValueError("shift ends before it starts (check End Date)")
            roster.append(RosterRec(r["Assignment ID"], cm, start, end, activities[aid], _yes(r.get("Is Locked", ""))))
        except (ValueError, KeyError) as exc:
            errors.append(f"roster_df line {n}: {exc}")

    # ---- demand slots ------------------------------------------------------
    slots: dict[tuple[str, datetime], dict] = {}
    n_inactive = n_excluded = 0
    for n, r in enumerate(tables["activity_demand_df"], start=2):
        try:
            aid = r.get("Activity ID") or by_name.get(r.get("Activity Name", ""), "")
            if aid not in activities:
                raise ValueError(f"unknown Activity ID {aid!r}")
            if not _yes(r.get("Is Active", "yes")):
                n_inactive += 1
                continue
            if aid in excluded:
                n_excluded += 1
                continue
            s_t, e_t = parse_shift(r["Shift"])
            units = int(r["Demand"])
            if units < 0:
                raise ValueError("Demand must be >= 0")
        except (ValueError, KeyError) as exc:
            errors.append(f"activity_demand_df line {n}: {exc}")
            continue
        for day in horizon_days:
            s, e = slot_datetimes(day, s_t, e_t)
            slot = slots.setdefault((aid, s), {"demand_id": r["Demand ID"], "end": e, "units": 0})
            slot["units"] += units
    if errors:
        raise DBDataError(errors)

    # ---- classify roster rows --------------------------------------------------
    report = Counter()
    disruptions: list[dict] = []
    regular_cover: Counter = Counter()
    fixed: dict[tuple[str, datetime], list[str]] = defaultdict(list)
    extra_fixed: list[RosterRec] = []  # locked Relief/Excess rows without a demand slot
    released: list[str] = []
    per_crew_rows: dict[str, list[RosterRec]] = defaultdict(list)
    for rec in roster:
        crew = crews[rec.cm_id]
        per_crew_rows[rec.cm_id].append(rec)
        in_horizon = h0 <= rec.start < h1
        if rec.activity.is_standby or not in_horizon:
            continue
        key = (rec.activity.activity_id, rec.start)
        if crew.crew_type is None:  # regular crew
            if key not in slots:
                report["regular_rows_without_demand"] += 1
                continue
            why = rules.reason(rec.cm_id, rec.activity.activity_id, rec.start, rec.end)
            if why:
                report["regular_rows_disrupted"] += 1
                disruptions.append({"assignment_id": rec.assignment_id, "cm_id": rec.cm_id,
                                    "activity": rec.activity.name, "start": rec.start.strftime("%Y-%m-%d %H:%M"),
                                    "reason": why})  # fmt: skip
            else:
                regular_cover[key] += 1
                report["regular_rows_covering"] += 1
            continue
        locked = rec.locked or not rec.activity.can_be_moved
        if not locked:
            released.append(rec.assignment_id)
            continue
        why = rules.reason(rec.cm_id, rec.activity.activity_id, rec.start, rec.end)
        if why:
            report["locked_rows_invalid"] += 1
            disruptions.append({"assignment_id": rec.assignment_id, "cm_id": rec.cm_id,
                                "activity": rec.activity.name, "start": rec.start.strftime("%Y-%m-%d %H:%M"),
                                "reason": why})  # fmt: skip
            continue
        if key in slots:
            fixed[key].append(rec.cm_id)
        else:
            extra_fixed.append(rec)

    # ---- Relief / Excess crew -----------------------------------------------
    crew_objs: dict[str, Crew] = {}
    excluded_crew: Counter = Counter()
    for cm, c in crews.items():
        ctype = c.crew_type
        if ctype is None:
            continue
        if not c.medically_fit:
            excluded_crew["medically_unfit"] += 1
            continue
        rows = sorted(per_crew_rows.get(cm, []), key=lambda x: x.start)
        runs = _runs(x.start.date() for x in rows)
        run = next((r for r in runs if r[-1] >= pc.horizon_start), None)
        if run is None or run[0] >= h1.date():
            excluded_crew["not_rostered_in_horizon"] += 1
            continue
        nxt_run = next((r for r in runs if r[0] > run[-1]), None)
        default = cfg.default_work_off_days.get(ctype.value, [len(run), 0])
        off = (nxt_run[0] - run[-1]).days - 1 if nxt_run else default[1]
        standby_starts = [x.start.time() for x in rows if x.activity.is_standby]
        rostered = Counter(standby_starts or [x.start.time() for x in rows]).most_common(1)[0][0]
        real_rows = [x for x in rows if not x.activity.is_standby]
        prior = [x for x in real_rows if x.start < h0]
        prior_rec = max(prior, key=lambda x: x.end) if prior else None
        current = None
        if prior_rec is not None and prior_rec.start.date() >= run[0] and prior_rec.start.time() != rostered:
            current = prior_rec.start.time()
        later = [x for x in real_rows if x.start >= h1]
        nxt = min(later, key=lambda x: x.start) if later else None
        # Location: onsite record for the start day (or the day before) > prior duty > depot.
        initial = c.depot
        if prior_rec is not None:
            initial = prior_rec.activity.end_depot
        for d, loc, _ in rules.onsite.get(cm, ()):
            if loc and pc.horizon_start - timedelta(days=1) <= d <= pc.horizon_start:
                initial = loc
        rostered_days = {x.start.date() for x in rows}
        unavailable = {d for d in horizon_days if rules.unavailable(cm, d)}
        if cfg.enforce_swing_work_days:
            unavailable |= {d for d in horizon_days if d not in rostered_days}
        crew_objs[cm] = Crew(
            crew_id=cm,
            crew_type=ctype,
            initial_location=initial,
            base_location=c.depot or None,
            rostered_start=rostered,
            current_start=current,
            swing_start=run[0],
            work_days=len(run),
            off_days=max(off, 0),
            prior_duty_end=prior_rec.end if prior_rec else None,
            prior_shift_type=_shift_of(prior_rec.start, cfg) if prior_rec else None,
            next_commitment_start=nxt.start if nxt else None,
            next_commitment_location=nxt.activity.start_depot if nxt else None,
            next_commitment_shift_type=_shift_of(nxt.start, cfg) if nxt else None,
            unavailable_dates=frozenset(unavailable),
        )

    # ---- duties ------------------------------------------------------------------
    duties: list[Duty] = []
    n_over = 0
    elig_sizes = []
    for (aid, start), slot in sorted(slots.items(), key=lambda kv: (kv[0][1], kv[0][0])):
        act = activities[aid]
        units = slot["units"]
        reg = regular_cover.get((aid, start), 0)
        assigned = tuple(cm for cm in fixed.get((aid, start), []) if cm in crew_objs)
        need = units - reg
        if need < len(assigned):
            n_over += 1
            need = len(assigned)
        if need <= 0:
            continue
        eligible = {cm for cm in crew_objs if rules.reason(cm, aid, start, slot["end"]) is None}
        eligible |= set(assigned)
        elig_sizes.append(len(eligible))
        duties.append(
            Duty(
                duty_id=f"{slot['demand_id']}-{start:%Y%m%d}",
                activity=act.name,
                start=start,
                end=slot["end"],
                start_location=act.start_depot,
                end_location=act.end_depot,
                demand=need,
                priority=ranks.get(act.head),
                eligible_crew=frozenset(eligible),
                assigned_crew=assigned,
            )
        )
    for rec in extra_fixed:
        if rec.cm_id not in crew_objs:
            continue
        duties.append(
            Duty(
                duty_id=f"R-{rec.assignment_id}",
                activity=rec.activity.name,
                start=rec.start,
                end=rec.end,
                start_location=rec.activity.start_depot,
                end_location=rec.activity.end_depot,
                demand=1,
                priority=ranks.get(rec.activity.head),
                eligible_crew=frozenset({rec.cm_id}),
                assigned_crew=(rec.cm_id,),
            )
        )

    total_units = sum(s["units"] for s in slots.values())
    prep = {
        "source_format": "optiflux_db_schema",
        "horizon_start": h0.strftime("%Y-%m-%d %H:%M"),
        "horizon_days": pc.horizon_days,
        "demand_slots": len(slots),
        "demand_units_total": total_units,
        "demand_rows_inactive": n_inactive,
        "demand_rows_excluded": n_excluded,
        "units_covered_by_regular_crew": sum(regular_cover.values()),
        "units_fixed_relief_excess": sum(len(d.assigned_crew) for d in duties),
        "open_demand_units": sum(d.open_demand for d in duties),
        "over_covered_slots": n_over,
        "roster_rows": len(roster),
        "released_unlocked_assignments": len(released),
        "regular_rows_disrupted": report["regular_rows_disrupted"],
        "regular_rows_without_demand": report["regular_rows_without_demand"],
        "locked_rows_invalid": report["locked_rows_invalid"],
        "disruptions_by_reason": dict(Counter(d["reason"] for d in disruptions)),
        "disruptions": disruptions[:500],
        "n_crew_master": len(crews),
        "n_crew_optimised": len(crew_objs),
        "n_crew_regular": sum(1 for c in crews.values() if c.crew_type is None),
        "crew_excluded_from_optimisation": dict(excluded_crew),
        "eligible_list_size": {
            "min": min(elig_sizes, default=0),
            "mean": round(sum(elig_sizes) / len(elig_sizes), 2) if elig_sizes else 0,
            "max": max(elig_sizes, default=0),
        },
        "activity_priority": ranks,
        "excluded_activities": sorted(excluded),
        "constraint_warnings": pc.warnings,
        "not_modelled": ["TEA Required", "Location (activity area; Start/End Depot are used)", "Position",
                         "Email/First Name/Last Name (identity only)"],  # fmt: skip
    }
    meta = dict(metadata or {})
    meta.update(model_config=cfg.to_dict(), data_preparation=prep)
    return Instance(
        name=meta.get("name", name),
        crews=crew_objs,
        duties=duties,
        horizon_start=h0,
        horizon_end=h1,
        qualifications=None,
        travel={},
        metadata=meta,
    )


def load_db_instance(directory: str | Path) -> Instance:
    """Load from a directory of table CSVs or from an Excel workbook (one sheet per table)."""

    base = Path(directory)
    if is_excel(base):
        return build_instance(read_excel_tables(base), name=base.stem, metadata={"name": base.stem})
    meta: dict = {}
    if (base / "instance.json").exists():
        meta = json.loads((base / "instance.json").read_text())
    meta.setdefault("name", base.name)
    return build_instance(read_tables(base), name=meta["name"], metadata=meta)
